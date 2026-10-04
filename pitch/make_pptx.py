"""PDF deck → editable-container PPTX (one full-slide image per page) + animated GIF demo slides.
usage: python pitch/make_pptx.py   (needs pymupdf + python-pptx)"""
import os, io, fitz
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
D = os.path.dirname(os.path.abspath(__file__))
pdf = fitz.open(os.path.join(D, "QAYDH_T0049_pitch.pdf"))
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
def add_img(stream_or_path):
    s = prs.slides.add_slide(blank); s.shapes.add_picture(stream_or_path, 0, 0, prs.slide_width, prs.slide_height); return s
for i, page in enumerate(pdf):
    png = page.get_pixmap(dpi=110).tobytes("png"); add_img(io.BytesIO(png))
    if i == 2:   # after the team slide: live demo GIFs
        for g, t in [("qaydh_musaffah_tour.gif", "Live demo · Musaffah tour"), ("qaydh_three_cities.gif", "Live demo · Riyadh, Abu Dhabi, Musaffah")]:
            s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0x14, 0x11, 0x0e)
            tb = s.shapes.add_textbox(Emu(400000), Emu(150000), Emu(11000000), Emu(500000)).text_frame; tb.text = t
            r = tb.paragraphs[0].runs[0]; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = RGBColor(0xff, 0x6b, 0x2c)
            s.shapes.add_picture(os.path.join(D, "gifs", g), Emu(600000), Emu(750000), width=Emu(10992000))
out = os.path.join(D, "QAYDH_T0049_pitch.pptx"); prs.save(out); print(out, f"{os.path.getsize(out)/1e6:.1f} MB", len(prs.slides), "slides")
